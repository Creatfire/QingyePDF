# SPDX-License-Identifier: AGPL-3.0-only
"""Offline document operations. Each invocation owns one isolated job directory."""
import csv
import base64
import contextlib
import difflib
import html
import io
import json
import math
import re
import sys
import zipfile
from pathlib import Path

import pymupdf as fitz

fitz.TOOLS.mupdf_display_errors(False)
fitz.TOOLS.mupdf_display_warnings(False)

def progress(done,total,label=''):
    print('QINGYE_PROGRESS '+json.dumps({'done':done,'total':total,'label':label},ensure_ascii=True),file=sys.stderr,flush=True)


def safe_tessdata(folder):
    # Tesseract opens traineddata with narrow fopen(); in the frozen build a
    # non-ASCII path (e.g. a Chinese install folder) cannot be opened at all.
    if not folder or str(folder).isascii():
        return folder
    import hashlib
    import shutil
    import tempfile
    key = hashlib.sha256(str(folder).encode('utf-8')).hexdigest()[:16]
    for base in (Path(tempfile.gettempdir()), Path('C:/ProgramData/QingyePDF')):
        if not str(base).isascii():
            continue
        try:
            target = base / ('qingye-tessdata-' + key)
            target.mkdir(parents=True, exist_ok=True)
            for source in Path(folder).glob('*.traineddata'):
                copy = target / source.name
                if not copy.is_file() or copy.stat().st_size != source.stat().st_size:
                    shutil.copyfile(source, copy)
            return target
        except OSError:
            continue
    raise ValueError('本地 OCR 语言目录路径无法被识别引擎读取，请把程序放到纯英文路径。')

def deskew_angle(image):
    # ponytail: projection profile detects modest text-page skew (±5°), not arbitrary photo perspective.
    sample=image.convert('L');sample.thumbnail((800,800));binary=sample.point(lambda x:0 if x<160 else 255).convert('1')
    bits=[i.bit_count() for i in range(256)]
    def score(angle):
        rotated=binary.rotate(angle,resample=0,expand=False,fillcolor=1);stride=(rotated.width+7)//8;data=rotated.tobytes();counts=[]
        for row in range(rotated.height):counts.append(sum(8-bits[b] for b in data[row*stride:(row+1)*stride]))
        average=sum(counts)/len(counts);return sum((n-average)**2 for n in counts)/len(counts)
    base=score(0);angle=max((step*.5 for step in range(-10,11)),key=score)
    return angle if abs(angle)>=.5 and score(angle)>base*1.03 else 0


def pages(spec, count):
    if not spec or str(spec).strip().lower() in ('all', '全部', '*'):
        return list(range(count))
    result = []
    for part in str(spec).replace('，', ',').split(','):
        match = re.fullmatch(r'\s*(\d+)(?:\s*-\s*(\d+))?\s*', part)
        if not match:
            raise ValueError('页码格式：1-3,5,8；留空表示全部页面。')
        first, last = int(match[1]), int(match[2] or match[1])
        if not 1 <= first <= count or not 1 <= last <= count:
            raise ValueError(f'页码应在 1–{count} 之间。')
        step = 1 if last >= first else -1
        result.extend(range(first - 1, last - 1 + step, step))
    return result


def number(value, low, high, default):
    value = float(default if value is None or value == '' else value)
    if not math.isfinite(value) or not low <= value <= high:
        raise ValueError(f'数值应在 {low}–{high} 之间。')
    return value


def color(value='#b91c1c'):
    if not re.fullmatch(r'#[0-9a-fA-F]{6}', value):
        raise ValueError('颜色格式无效。')
    return tuple(int(value[i:i + 2], 16) / 255 for i in (1, 3, 5))


def region(page, values):
    values = values or [0.15, 0.15, 0.75, 0.3]
    if len(values) != 4:
        raise ValueError('请框选一个有效区域。')
    x0, y0, x1, y1 = [number(v, 0, 1, 0) for v in values]
    if x1 - x0 < 0.002 or y1 - y0 < 0.002:
        raise ValueError('所选区域太小。')
    r = page.rect
    return fitz.Rect(x0 * r.width, y0 * r.height, x1 * r.width, y1 * r.height) * page.derotation_matrix


def open_pdf(file, password=''):
    doc = fitz.open(file)
    if doc.needs_pass and not doc.authenticate(password):
        doc.close()
        raise ValueError('PDF 密码不正确，请输入该文件的密码。')
    return doc


def save(doc, target, **options):
    doc.save(target, garbage=4, deflate=True, use_objstms=1,
             encryption=options.pop('encryption', fitz.PDF_ENCRYPT_NONE), **options)


def text_box(page, rect, text, request, border=False):
    size = number(request.get('size'), 6, 100, 18)
    opacity = number(request.get('opacity'), 0.05, 1, 1)
    ink = request.get('color', '#b91c1c')
    color(ink)
    if border:
        page.draw_rect(rect, color=color(ink), width=2, stroke_opacity=opacity)
        rect = rect + (6, 6, -6, -6)
    # MuPDF's HTML renderer includes fallback CJK fonts; no system font dependency.
    css = f'body {{font-family:sans-serif;font-size:{size}pt;color:{ink};margin:0;}}'
    spare, scale = page.insert_htmlbox(rect, '<div>' + html.escape(text).replace('\n', '<br>') + '</div>', css=css, opacity=opacity)
    if spare < 0:
        raise ValueError('文字无法放入所选区域，请扩大区域或缩小字号。')


def sharpen_images(doc, request):
    from PIL import Image, ImageFilter
    levels = {'mild': (1.0, 80, 2), 'standard': (1.5, 150, 3), 'strong': (2.0, 220, 3)}
    level = request.get('strength', 'standard')
    if level not in levels:
        raise ValueError('请选择轻度、标准或较强锐化。')
    images = {}
    # An image object can be shared by several pages, so sharpen the whole
    # document and replace each embedded image once, preserving page structure.
    for index, page in enumerate(doc):
        for item in page.get_images(full=True):
            images.setdefault(item[0], (index, item[1], item[2], item[3]))
    for step, (xref, (index, mask, width, height)) in enumerate(images.items()):
        progress(step, len(images), '锐化 PDF 图像')
        if width * height > 40_000_000:
            raise ValueError('嵌入图像超过 4000 万像素，请先降低图像尺寸。')
        pix = fitz.Pixmap(doc, xref)
        if mask:
            pix = fitz.Pixmap(pix, fitz.Pixmap(doc, mask))
        if pix.colorspace.n != 3:
            pix = fitz.Pixmap(fitz.csRGB, pix)
        image = Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGBA' if pix.alpha else 'RGB')
        alpha = image.getchannel('A') if pix.alpha else None
        enhanced = image.convert('RGB').filter(ImageFilter.UnsharpMask(*levels[level]))
        if alpha is not None:
            enhanced.putalpha(alpha)
        buffer = io.BytesIO()
        enhanced.save(buffer, format='PNG')
        doc[index].replace_image(xref, stream=buffer.getvalue())
    progress(len(images), len(images), '锐化完成')
    return len(images)


def add_ocr_layer(doc, selected, request):
    tessdata = request.get('tessdata')
    if not tessdata or not (Path(tessdata) / 'chi_sim.traineddata').is_file():
        raise ValueError('本地 OCR 语言资源缺失。')
    recognized = 0
    for step, index in enumerate(selected):
        progress(step, len(selected), '识别扫描文字')
        page = doc[index]
        if page.get_text().strip():
            continue
        if page.rect.width * page.rect.height * (200 / 72) ** 2 > 40_000_000:
            raise ValueError('扫描页超过 4000 万像素，请先降低页面尺寸。')
        rotation = page.rotation
        with contextlib.ExitStack() as stack:
            candidates = []
            for display_rotation in dict.fromkeys((0, rotation)):
                page.set_rotation(display_rotation)
                try:
                    pix = page.get_pixmap(dpi=200, alpha=False, annots=False)
                finally:
                    page.set_rotation(rotation)
                layer = stack.enter_context(fitz.open(stream=pix.pdfocr_tobytes(language='chi_sim+eng', tessdata=tessdata), filetype='pdf'))
                text = layer[0].get_text().strip()
                # ponytail: compare coherent word lengths for raw/display
                # orientations; this is a fallback heuristic, not confidence.
                score = sum(len(word) ** 2 for word in re.findall(r'[A-Za-z0-9\u4e00-\u9fff]{3,}', text))
                candidates.append((score, layer, display_rotation, text))
            _, layer, text_rotation, text = max(candidates, key=lambda item: item[0])
            if not text:
                continue
            # Transfer only invisible OCR text. The original scan, vectors,
            # annotations, widgets and intrinsic rotation remain in the PDF.
            for image in layer[0].get_images():
                layer[0].delete_image(image[0])
            page.show_pdf_page(page.rect * page.derotation_matrix, layer, 0, rotate=text_rotation)
            recognized += 1
    progress(len(selected), len(selected), '文字识别完成')
    return recognized


def export(doc, request, output):
    selected = pages(request.get('pages'), len(doc))
    fmt = request.get('format', 'png')
    target = output / ('result.' + ('zip' if fmt in ('png', 'jpg', 'svg') else fmt))
    note = ''
    if fmt in ('png', 'jpg', 'svg'):
        dpi = number(request.get('dpi'), 36, 300, 144)
        with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as z:
            for step,index in enumerate(selected):
                progress(step,len(selected),'导出页面')
                page = doc[index]
                if fmt == 'svg':
                    data = page.get_svg_image().encode('utf-8')
                else:
                    if page.rect.width * page.rect.height * (dpi / 72) ** 2 > 40_000_000:
                        raise ValueError('单页图像超过 4000 万像素，请降低 DPI。')
                    pix = page.get_pixmap(dpi=int(dpi), alpha=False)
                    data = pix.tobytes('jpeg' if fmt == 'jpg' else 'png')
                z.writestr(f'page-{index + 1:05d}.{fmt}', data)
    elif fmt in ('txt', 'html'):
        if fmt == 'txt':
            text = '\n\f\n'.join(doc[i].get_text(sort=True) for i in selected)
        else:
            text = '<!doctype html><meta charset="UTF-8"><title>PDF 导出</title><style>body{background:#ddd}.page{margin:20px auto;background:white}</style>'
            text += ''.join(doc[i].get_text('html') for i in selected)
        target.write_text(text, encoding='utf-8')
    elif fmt == 'docx':
        from docx import Document
        from docx.shared import Inches
        word = Document()
        for n, i in enumerate(selected):
            progress(n,len(selected),'转换 Word')
            if n:
                word.add_page_break()
            page = doc[i]
            for block in page.get_text('blocks', sort=True):
                if block[6] == 0:
                    word.add_paragraph(block[4].strip())
            for image in page.get_images(full=True):
                extracted = doc.extract_image(image[0])
                try:
                    word.add_picture(io.BytesIO(extracted['image']), width=Inches(5.5))
                except Exception:
                    pass  # Unsupported image encodings do not prevent text conversion.
        word.save(target)
        note = 'Word 为可编辑文本与图片转换；复杂版式、公式和分栏不保证保真。'
    elif fmt in ('xlsx', 'csv'):
        from openpyxl import Workbook
        book = Workbook()
        book.remove(book.active)
        rows_all = []
        for i in selected:
            progress(selected.index(i),len(selected),'转换表格')
            rows = []
            for table in doc[i].find_tables().tables:
                rows.extend(table.extract())
                rows.append([])
            if not rows:
                rows = [[line] for line in doc[i].get_text(sort=True).splitlines()]
            sheet = book.create_sheet(f'Page {i + 1}')
            for row in rows:
                # Prevent spreadsheet formula execution from document text.
                safe = ["'" + cell if isinstance(cell, str) and cell.startswith(('=', '+', '-', '@')) else cell for cell in row]
                sheet.append(safe)
                rows_all.append(safe)
        if fmt == 'xlsx':
            book.save(target)
        else:
            with target.open('w', encoding='utf-8-sig', newline='') as file:
                csv.writer(file).writerows(rows_all)
        note = '优先识别表格；未识别到表格时逐行导出文本。'
    elif fmt == 'pptx':
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        first = doc[selected[0]].rect
        deck.slide_width = Inches(first.width / 72)
        deck.slide_height = Inches(first.height / 72)
        for i in selected:
            progress(selected.index(i),len(selected),'转换幻灯片')
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            pix = doc[i].get_pixmap(dpi=144, alpha=False)
            slide.shapes.add_picture(io.BytesIO(pix.tobytes('png')), 0, 0, width=deck.slide_width, height=deck.slide_height)
        deck.save(target)
        note = 'PowerPoint 每页转换为图像幻灯片，保持视觉版式，文字不能单独编辑。'
    else:
        raise ValueError('不支持的导出格式。')
    return {'files': [target.name], 'note': note}


def import_document(file, output):
    extension = file.suffix.lower()
    if extension in ('.png', '.jpg', '.jpeg', '.bmp', '.tif', '.tiff', '.webp'):
        from PIL import Image, ImageSequence
        pdf = fitz.open()
        with Image.open(file) as image:
            for frame in ImageSequence.Iterator(image):
                bitmap = frame.convert('RGB')
                buffer = io.BytesIO()
                bitmap.save(buffer, format='PNG')
                page = pdf.new_page(width=bitmap.width * .75, height=bitmap.height * .75)
                page.insert_image(page.rect, stream=buffer.getvalue())
        save(pdf, output / 'result.pdf')
        pdf.close()
        return {'files': ['result.pdf'], 'note': '图像已转换为 PDF。'}
    if extension not in ('.docx', '.xlsx', '.pptx', '.txt', '.html', '.htm', '.epub', '.xps', '.cbz', '.svg'):
        raise ValueError('支持现代 Office 格式 DOCX/XLSX/PPTX、图像、TXT、HTML、EPUB、XPS、CBZ 与 SVG。')
    # Standard open-source MuPDF imports Office as reflowed HTML, not its commercial Pro renderer.
    if extension in ('.txt', '.html', '.htm'):
        content = file.read_text(encoding='utf-8-sig')
        if extension == '.txt':
            content = '<pre style="white-space:pre-wrap">' + html.escape(content) + '</pre>'
        # Restrict HTML to text when importing; never fetch external resources.
        if extension in ('.html', '.htm'):
            from lxml import html as lh
            content = '<div>' + html.escape(lh.fromstring(content).text_content()).replace('\n', '<br>') + '</div>'
        stream = io.BytesIO()
        writer = fitz.DocumentWriter(stream)
        story = fitz.Story(html=content)
        more = True
        while more:
            device = writer.begin_page(fitz.Rect(0, 0, 595, 842))
            more, _ = story.place(fitz.Rect(40, 40, 555, 802))
            story.draw(device)
            writer.end_page()
        writer.close()
        (output / 'result.pdf').write_bytes(stream.getvalue())
    else:
        with fitz.open(file) as source:
            (output / 'result.pdf').write_bytes(source.convert_to_pdf())
    return {'files': ['result.pdf'], 'note': 'Office 输入采用内容重排，原分页和复杂版式可能变化。' if extension in ('.docx', '.xlsx', '.pptx') else ''}


def process(request, input_file, output):
    output.mkdir(parents=True, exist_ok=True)
    action = request.get('action')
    if action in ('ocr', 'ocr-layer', 'scan'):
        relocated = safe_tessdata(request.get('tessdata'))
        if relocated is not None:
            request = {**request, 'tessdata': str(relocated)}
    progress(0,1,action or '本地处理')
    if action == 'import':
        return import_document(input_file, output)
    with open_pdf(input_file, request.get('password', '')) as doc:
        encrypted=doc.xref_get_key(-1,'Encrypt')[0]!='null'
        owner=bool(doc.authenticate(request.get('password',''))&4) if encrypted else True
        if encrypted and not owner:
            if action in ('export','notes-export','compare','ocr','ocr-layer','scan') and not doc.permissions&fitz.PDF_PERM_COPY:
                raise ValueError('该 PDF 不允许复制 / 转换内容，请输入管理密码。')
            if action not in ('inspect','export','notes-export','compare','decrypt') and not doc.permissions&fitz.PDF_PERM_MODIFY:
                raise ValueError('该 PDF 不允许编辑，请输入管理密码。')
        if action == 'inspect':
            index = int(number(request.get('page'), 1, len(doc), 1)) - 1
            page = doc[index]
            blocks = []
            readable=not encrypted or owner or bool(doc.permissions&fitz.PDF_PERM_COPY)
            for block in page.get_text('blocks', sort=True) if readable else []:
                if block[6] == 0:
                    r = fitz.Rect(block[:4]) * page.rotation_matrix
                    blocks.append({'text': block[4], 'rect': [r.x0 / page.rect.width, r.y0 / page.rect.height, r.x1 / page.rect.width, r.y1 / page.rect.height]})
            data={'toc': doc.get_toc(), 'metadata': doc.metadata, 'pages': len(doc), 'blocks': blocks, 'encrypted': encrypted}
            if request.get('imageRects'):
                data['images']=[list(fitz.Rect(image['bbox'])*~page.transformation_matrix) for image in page.get_image_info()][:2000]
                if request.get('imageRects')=='all':
                    data['imagePages']=[[list(fitz.Rect(image['bbox'])*~p.transformation_matrix) for image in p.get_image_info()][:2000] for p in doc]
            if request.get('annotations') and readable:
                notes=[]
                for n,p in enumerate(doc):
                    for a in p.annots() or []:
                        if a.type[1] in ('Popup','Link'):continue
                        stroke=a.colors.get('stroke') or a.colors.get('fill') or []
                        tint='#'+''.join(f'{max(0,min(255,round(c*255))):02x}' for c in stroke[:3]) if len(stroke)>=3 else ''
                        excerpt=p.get_textbox(a.rect).strip() if a.type[1] in ('Highlight','Underline','StrikeOut','Squiggly') else ''
                        notes.append({'page':n+1,'type':a.type[1],'text':a.info.get('content',''),'excerpt':excerpt,'color':tint,'rect':list(a.rect*~p.transformation_matrix),'xref':a.xref})
                data['notes']=notes
            if request.get('headings') and readable:
                lines=[];sizes={}
                for n,p in enumerate(doc):
                    for b in p.get_text('dict',flags=fitz.TEXTFLAGS_DICT & ~fitz.TEXT_PRESERVE_IMAGES)['blocks']:
                        for line in b.get('lines',[]):
                            spans=line['spans'];text=''.join(s['text'] for s in spans).strip()
                            if not text:continue
                            size=round(max(s['size'] for s in spans),1);sizes[size]=sizes.get(size,0)+len(text);lines.append((n+1,text,size))
                body=max(sizes,key=sizes.get) if sizes else 12
                candidates=[(n,t,z) for n,t,z in lines if len(t)<=100 and (z>=body*1.18 or re.match(r'^(第[一二三四五六七八九十百\d]+[章节篇部]|\d+(?:\.\d+){0,3}[、.\s])',t))]
                levels=sorted({z for _,_,z in candidates},reverse=True)[:5];headings=[];previous=0
                for n,t,z in candidates[:1000]:
                    level=min((levels.index(z)+1 if z in levels else len(levels) or 1),previous+1)
                    if headings and headings[-1][1:]==[t,n]:continue
                    headings.append([level,t,n]);previous=level
                data['headings']=headings
            return {'data':data,'files':[]}
        if action == 'notes-export':
            from docx import Document
            word=Document();word.add_heading('PDF 批注摘录',0)
            for note in request.get('notes',[]):
                word.add_heading(f"第 {note['page']} 页 · {note['type']}",2)
                if note.get('excerpt'):word.add_paragraph(note['excerpt'],style='Quote')
                if note.get('text'):word.add_paragraph(note['text'])
                if note.get('color'):word.add_paragraph('颜色：'+note['color'])
            word.save(output/'result.docx');return {'files':['result.docx']}
        if action == 'export':
            return export(doc, request, output)
        if action == 'compare':
            sources=request.get('inputs', [])
            if len(sources)!=1:
                raise ValueError('比较需要另一个 PDF。')
            with open_pdf(Path(sources[0]),request.get('password','')) as other:
                if other.xref_get_key(-1,'Encrypt')[0]!='null' and not other.authenticate(request.get('password',''))&4 and not other.permissions&fitz.PDF_PERM_COPY:raise ValueError('比较文档不允许复制，请输入管理密码。')
                chunks=['<!doctype html><meta charset="UTF-8"><title>PDF 文字差异</title><style>table{width:100%;font:13px monospace}.diff_add{background:#bbefcc}.diff_sub{background:#f7bdc1}.diff_chg{background:#f7e5a2}td{white-space:pre-wrap}body{padding:20px}</style><h1>逐页文字差异报告</h1><p>仅比较提取的文字；图像与版式差异请并排检查。</p>']
                offset=int(number(request.get('offset'),-100000,100000,0));visual=request.get('visual',False)
                if abs(offset)>=max(len(doc),len(other)):raise ValueError('页差过大，两份文档没有可比较的页面。')
                if visual:chunks=['<!doctype html><meta charset="UTF-8"><title>PDF 图像差异</title><style>body{font:14px sans-serif;padding:20px}img{max-width:32%;vertical-align:top;border:1px solid #aaa}</style><h1>页面图像差异报告</h1><p>96 DPI 渲染、左上角对齐；红色显示像素差异。字体抗锯齿、版式移动和尺寸变化也会产生差异。</p>']
                indexes=pages(request.get('pages'),len(doc)) if request.get('pages') else list(range(max(len(doc),len(other)-offset)))
                if visual and len(indexes)>200:raise ValueError('图像比较一次最多 200 页，请指定页码范围。')
                for step,i in enumerate(indexes):
                    progress(step,len(indexes),'比较页面');j=i+offset
                    left=doc[i].get_text(sort=True).splitlines() if i<len(doc) else []
                    right=other[j].get_text(sort=True).splitlines() if 0<=j<len(other) else []
                    if not visual:chunks.append(f'<h2>当前第 {i+1} 页 / 对比第 {j+1} 页</h2>'+difflib.HtmlDiff(wrapcolumn=80).make_table(left,right,fromdesc='当前文档',todesc='比较文档',context=True,numlines=3));continue
                    from PIL import Image,ImageChops
                    images=[]
                    for source,pno in [(doc,i),(other,j)]:
                        if 0<=pno<len(source):
                            p=source[pno]
                            if p.rect.width*p.rect.height*(96/72)**2>15_000_000:raise ValueError('比较页面超过 1500 万像素。')
                            images.append(Image.open(io.BytesIO(p.get_pixmap(dpi=96,alpha=False).tobytes('png'))).convert('RGB'))
                        else:images.append(None)
                    size=(max(im.width for im in images if im),max(im.height for im in images if im));aligned=[]
                    for im in images:
                        canvas=Image.new('RGB',size,'white')
                        if im:canvas.paste(im,(0,0))
                        aligned.append(canvas)
                    diff=ImageChops.difference(*aligned).convert('L').point(lambda v:255 if v>24 else 0);marked=Image.composite(Image.new('RGB',size,'#ff5555'),aligned[0],diff)
                    count=diff.histogram()[255];chunks.append(f'<h2>当前第 {i+1} 页 / 对比第 {j+1} 页 · 差异像素 {count}</h2>')
                    for im in [*aligned,marked]:
                        buffer=io.BytesIO();im.save(buffer,format='PNG');chunks.append('<img src="data:image/png;base64,'+base64.b64encode(buffer.getvalue()).decode()+'">')
                (output/'result.html').write_text(''.join(chunks),encoding='utf-8')
            return {'files':['result.html'],'note':'已导出图像差异标记报告。' if visual else '已导出逐页文字差异报告。'}
        selected = pages(request.get('pages'), len(doc))
        target = output / 'result.pdf'
        note = ''
        if action == 'encrypt':
            user = str(request.get('userPassword', ''))
            owner = str(request.get('ownerPassword', ''))
            if not owner or len(user.encode()) > 127 or len(owner.encode()) > 127:
                raise ValueError('必须设置管理密码，密码最多为 127 个 UTF-8 字节。')
            if owner == user:
                raise ValueError('管理密码应与打开密码不同。')
            permissions = fitz.PDF_PERM_ACCESSIBILITY
            for name, flag in [('print', fitz.PDF_PERM_PRINT | fitz.PDF_PERM_PRINT_HQ), ('copy', fitz.PDF_PERM_COPY), ('edit', fitz.PDF_PERM_MODIFY | fitz.PDF_PERM_ANNOTATE | fitz.PDF_PERM_FORM)]:
                if request.get(name, True):
                    permissions |= flag
            save(doc, target, encryption=fitz.PDF_ENCRYPT_AES_256, owner_pw=owner, user_pw=user, permissions=permissions)
            return {'files': ['result.pdf'], 'note': '已使用 AES-256 加密；权限限制依赖阅读器遵守。'}
        elif action == 'decrypt':
            # A supplied open password is insufficient to remove owner restrictions.
            if encrypted and not owner:
                raise ValueError('移除权限限制需要该 PDF 的管理密码。')
        elif action == 'organize':
            mode = request.get('mode', 'extract')
            if mode in ('extract', 'reorder'):
                if len(set(selected))!=len(selected):
                    result=fitz.open()
                    for step,index in enumerate(selected):
                        progress(step,len(selected),'复制页面');result.insert_pdf(doc,from_page=index,to_page=index,final=step==len(selected)-1)
                    result.set_metadata({k:v for k,v in doc.metadata.items() if isinstance(v,str) and v and k not in ('format','encryption')})
                    toc=[];previous=0
                    for level,title,page in doc.get_toc():
                        if page-1 not in selected:continue
                        level=min(level,previous+1);toc.append([level,title,selected.index(page-1)+1]);previous=level
                    result.set_toc(toc);save(result,target);result.close();return {'files':['result.pdf']}
                doc.select(selected)
            elif mode == 'delete':
                if len(set(selected)) == len(doc):
                    raise ValueError('至少保留一页。')
                doc.delete_pages(sorted(set(selected)))
            elif mode == 'rotate':
                angle = int(number(request.get('angle'), -270, 270, 90))
                if angle % 90:
                    raise ValueError('旋转角度应为 90 的倍数。')
                for i in selected:
                    doc[i].set_rotation((doc[i].rotation + angle) % 360)
            elif mode == 'blank':
                doc.new_page(pno=int(number(request.get('position'), 0, len(doc), len(doc))))
            elif mode == 'split':
                with zipfile.ZipFile(output / 'result.zip', 'w', zipfile.ZIP_DEFLATED) as z:
                    for i in selected:
                        piece = fitz.open()
                        piece.insert_pdf(doc, from_page=i, to_page=i)
                        z.writestr(f'page-{i + 1:05d}.pdf', piece.tobytes(garbage=4, deflate=True))
                        piece.close()
                return {'files': ['result.zip'], 'note': '已将每页分别拆为 PDF。'}
            elif mode == 'merge':
                for source in request.get('inputs', []):
                    with open_pdf(Path(source), request.get('mergePassword', '')) as other:
                        doc.insert_pdf(other)
            else:
                raise ValueError('不支持的页面操作。')
        elif action == 'outline':
            toc = request.get('toc', [])
            if not isinstance(toc, list) or len(toc) > 10000:
                raise ValueError('大纲条目无效或过多。')
            validated = []
            previous = 0
            for row in toc:
                level = int(number(row[0], 1, 32, 1))
                if level > previous + 1:
                    raise ValueError('大纲层级不能跳级，首条必须是一级。')
                title = str(row[1]).strip()[:500]
                if not title:
                    raise ValueError('大纲标题不能为空。')
                validated.append([level, title, int(number(row[2], 1, len(doc), 1))])
                previous = level
            doc.set_toc(validated)
        elif action == 'compress':
            if request.get('lossy'):
                dpi = int(number(request.get('dpi'), 50, 300, 144))
                quality = int(number(request.get('quality'), 10, 100, 65))
                doc.rewrite_images(dpi_threshold=dpi + 20, dpi_target=dpi, quality=quality)
                note = '图像已重压缩，清晰度可能下降。'
            doc.subset_fonts()
        elif action == 'sharpen':
            count = sharpen_images(doc, request)
            if not count:
                return {'files': [], 'unchanged': True, 'note': '没有可锐化的嵌入图像；矢量文字和图形保持原样。'}
            note = f'已锐化 {count} 个嵌入图像；保留文字、页面和批注。'
        elif action == 'ocr-layer':
            count = add_ocr_layer(doc, selected, request)
            if not count:
                return {'files': [], 'unchanged': True, 'note': '本页已有文字层，或没有识别到可用文字。'}
            note = f'已为 {count} 页添加可选择的中英文文字层；原页面外观保留，识别内容请复核。'
        elif action == 'flatten':
            doc.bake(annots=True, widgets=True)
        elif action == 'scan':
            from PIL import Image,ImageOps
            result=fitz.open();tessdata=request.get('tessdata')
            for index,p in enumerate(doc):
                progress(index,len(doc),'优化扫描页')
                if index not in selected:result.insert_pdf(doc,from_page=index,to_page=index);continue
                if p.rect.width*p.rect.height*(200/72)**2>40_000_000:raise ValueError('扫描页超过 4000 万像素，请先裁剪或降低页面尺寸。')
                pix=p.get_pixmap(dpi=200,alpha=False);image=Image.open(io.BytesIO(pix.tobytes('png'))).convert('RGB')
                angle=int(number(request.get('angle'),0,270,0))
                if angle%90:raise ValueError('扫描旋转角度应为 90 的倍数。')
                if angle:image=image.rotate(-angle,expand=True,fillcolor='white')
                if request.get('deskew'):image=image.rotate(deskew_angle(image),resample=Image.Resampling.BICUBIC,expand=False,fillcolor='white')
                if request.get('clean'):image=ImageOps.autocontrast(image,cutoff=1);image=image.point(lambda v:255 if v>int(number(request.get('whiteThreshold'),180,250,235)) else v)
                if request.get('gray'):image=image.convert('L')
                buffer=io.BytesIO();image.save(buffer,format='PNG')
                if request.get('recognize'):
                    ocr_image=fitz.Pixmap(buffer.getvalue())
                    if ocr_image.n!=3:ocr_image=fitz.Pixmap(fitz.csRGB,ocr_image)
                    ocr_image.set_dpi(200,200)
                    with fitz.open(stream=ocr_image.pdfocr_tobytes(language='chi_sim+eng',tessdata=tessdata),filetype='pdf') as piece:result.insert_pdf(piece)
                else:
                    page=result.new_page(width=image.width*72/200,height=image.height*72/200);page.insert_image(page.rect,stream=buffer.getvalue())
            result.set_toc(doc.get_toc());save(result,target);result.close();return {'files':['result.pdf'],'note':'已按 200 DPI 重新生成扫描页；原页面编辑对象与批注外观已栅格化，请检查纠偏及 OCR 结果。'}
        elif action == 'ocr':
            tessdata = request.get('tessdata')
            if not tessdata or not (Path(tessdata) / 'chi_sim.traineddata').is_file():
                raise ValueError('本地 OCR 语言资源缺失。')
            result = fitz.open()
            for index in range(len(doc)):
                progress(index,len(doc),'中英文 OCR')
                page = doc[index]
                if index not in selected or (page.get_text().strip() and not request.get('force')):
                    result.insert_pdf(doc, from_page=index, to_page=index)
                else:
                    pix = page.get_pixmap(dpi=200, alpha=False)
                    with fitz.open(stream=pix.pdfocr_tobytes(language='chi_sim+eng', tessdata=tessdata), filetype='pdf') as scanned:
                        result.insert_pdf(scanned)
            result.set_toc(doc.get_toc())
            save(result, target)
            result.close()
            return {'files': ['result.pdf'], 'note': '已生成带中英文文字层的 PDF；识别页按 200 DPI 图像输出，识别内容请复核。'}
        elif action in ('redact', 'text', 'stamp', 'image', 'page-stamp', 'number', 'watermark', 'shape', 'annotation', 'form', 'crop'):
            if action == 'redact':
                doc.bake(annots=True, widgets=True)
            for i in selected:
                page = doc[i]
                rect = region(page, request.get('rect'))
                ink = color(request.get('color', '#b91c1c'))
                text = str(request.get('text', ''))[:10000]
                if action in ('redact', 'text'):
                    page.add_redact_annot(rect, fill=(0, 0, 0) if action == 'redact' else (1, 1, 1))
                    page.apply_redactions(images=2 if action == 'redact' else 0, graphics=2 if action == 'redact' else 0, text=0)
                    if action == 'text':
                        text_box(page, rect, text, request)
                elif action in ('stamp', 'watermark', 'text', 'number'):
                    if action == 'number':
                        text = text.replace('{page}', str(i + 1)).replace('{total}', str(len(doc)))
                    text_box(page, rect, text, request, border=action == 'stamp')
                elif action == 'image':
                    file = request.get('asset')
                    if not file:
                        raise ValueError('请选择图片印章。')
                    # MuPDF keeps proportions on its own; the explicit centered fit
                    # keeps the placement independent of MuPDF's internal choice,
                    # and keep_proportion=False gives the deliberate stretch mode.
                    keep = bool(request.get('keepRatio', True))
                    placed = rect
                    if keep:
                        from PIL import Image
                        with Image.open(file) as picture:
                            ratio = picture.width / picture.height
                        if rect.width / rect.height > ratio:
                            width = rect.height * ratio
                            placed = fitz.Rect(rect.x0 + (rect.width - width) / 2, rect.y0, rect.x0 + (rect.width + width) / 2, rect.y1)
                        else:
                            height = rect.width / ratio
                            placed = fitz.Rect(rect.x0, rect.y0 + (rect.height - height) / 2, rect.x1, rect.y0 + (rect.height + height) / 2)
                    page.insert_image(placed, filename=file, keep_proportion=keep)
                elif action == 'page-stamp':
                    sources=request.get('inputs',[])
                    if len(sources)!=1:
                        raise ValueError('页面印章需要选择一个 PDF 来源。')
                    with open_pdf(Path(sources[0]),request.get('password','')) as source:
                        source.bake(annots=True,widgets=True)
                        pno=int(number(request.get('stampPage'),1,len(source),1))-1
                        page.show_pdf_page(rect,source,pno=pno)
                elif action == 'shape':
                    mode = request.get('shape', 'rect')
                    draw = {'rect': page.draw_rect, 'ellipse': page.draw_oval}.get(mode)
                    if mode == 'line':
                        page.draw_line(rect.top_left, rect.bottom_right, color=ink, width=2)
                    elif draw:
                        draw(rect, color=ink, width=2, fill=ink if request.get('fill') else None, fill_opacity=.2)
                    else:
                        raise ValueError('不支持的图形。')
                elif action == 'annotation':
                    method = {'highlight': page.add_highlight_annot, 'underline': page.add_underline_annot, 'strikeout': page.add_strikeout_annot}.get(request.get('annotation', 'underline'))
                    if not method:
                        raise ValueError('不支持的批注。')
                    annot = method(rect)
                    annot.set_colors(stroke=ink)
                    annot.set_info(content=text)
                    annot.update()
                elif action == 'form':
                    widget = fitz.Widget()
                    widget.field_name = (text or 'field') + f'_{i + 1}'
                    widget.field_type = fitz.PDF_WIDGET_TYPE_CHECKBOX if request.get('fieldType') == 'checkbox' else fitz.PDF_WIDGET_TYPE_TEXT
                    widget.rect = rect
                    widget.field_value = ''
                    widget.border_width = 1
                    page.add_widget(widget)
                elif action == 'crop':
                    page.set_cropbox(rect)
            if action == 'redact':
                doc.scrub(clean_pages=True, hidden_text=True, redactions=False, remove_links=False, reset_fields=False)
                note = '框选区域的文字、图像与相交图形已删除；同时清理元数据、附件与脚本。请检查导出副本。'
        else:
            raise ValueError('不支持的文档操作。')
        save(doc, target)
        if action == 'compress' and target.stat().st_size >= input_file.stat().st_size:
            target.write_bytes(input_file.read_bytes())
            note = '未获得更小文件，已保留原始数据与质量。'
        return {'files': ['result.pdf'], 'note': note}


if __name__ == '__main__':
    try:
        job = json.loads(Path(sys.argv[1]).read_text(encoding='utf-8-sig'))
        with contextlib.redirect_stdout(sys.stderr):
            result = process(job['request'], Path(job['input']), Path(job['output']))
        progress(1,1,'完成')
        print(json.dumps({'ok': True, **result}, ensure_ascii=True))
    except Exception as error:
        print(json.dumps({'ok': False, 'error': str(error)}, ensure_ascii=True))
        sys.exit(1)
