"""Run with the build virtualenv: python backend/test_worker.py."""
import io
import sys
import tempfile
import zipfile
from pathlib import Path
import pymupdf as fitz
from worker import process, pages


def check():
    root=Path(__file__).resolve().parents[1]
    with tempfile.TemporaryDirectory() as folder:
        base=Path(folder)
        source=base/'input.pdf'
        with fitz.open() as doc:
            page=doc.new_page(width=400,height=500)
            page.insert_text((50,80),'SECRET ORIGINAL',fontsize=24)
            page.insert_text((50,180),'Keep this content',fontsize=18)
            page.add_text_annot((30,300),'Saved comment')
            doc.new_page(width=400,height=500).insert_text((50,80),'Second page')
            doc.set_toc([[1,'First',1],[2,'Second',2]])
            doc.save(source)
        counter=0
        def run(request,input_file=source):
            nonlocal counter
            counter+=1
            output=base/f'job-{counter}'
            result=process(request,input_file,output)
            return result,output
        assert pages('2,1-2,2-1',2)==[1,0,1,1,0]
        try:pages('3',2);assert False
        except ValueError:pass
        info,_=run({'action':'inspect'})
        assert info['data']['toc'][0][1]=='First'
        assert any('SECRET' in b['text'] for b in info['data']['blocks'])
        detail,_=run({'action':'inspect','annotations':True,'headings':True,'imageRects':'all'})
        assert detail['data']['notes'][0]['text']=='Saved comment'
        assert any('SECRET' in row[1] for row in detail['data']['headings'])
        assert len(detail['data']['imagePages'])==2
        _,duplicates=run({'action':'organize','mode':'reorder','pages':'2,1,2,1'})
        with fitz.open(duplicates/'result.pdf') as d:
            assert len(d)==4 and 'Second page' in d[0].get_text() and 'SECRET' in d[1].get_text()
            assert any(a.info['content']=='Saved comment' for a in d[1].annots())
        _,notes=run({'action':'notes-export','notes':[{'page':1,'type':'Text','text':'中文批注','excerpt':'Selected text','color':'#ff0000'}]})
        from docx import Document
        assert '中文批注' in '\n'.join(p.text for p in Document(notes/'result.docx').paragraphs)
        _,visual=run({'action':'compare','inputs':[str(source)],'visual':True})
        report=(visual/'result.html').read_text(encoding='utf-8');assert 'data:image/png;base64,' in report and '差异像素 0' in report
        _,scan=run({'action':'scan','pages':'1','deskew':True,'gray':True,'clean':True,'recognize':True,'tessdata':str(root/'vendor'/'ocr')})
        with fitz.open(scan/'result.pdf') as d:
            assert len(d)==2 and abs(d[0].rect.width-400)<2
            assert 'SECRET' in d[0].get_text()
        for fmt in ('png','jpg','svg','html','txt','docx','xlsx','csv','pptx'):
            result,out=run({'action':'export','format':fmt})
            file=out/result['files'][0]
            assert file.stat().st_size>0
            if fmt in ('png','jpg','svg'):
                with zipfile.ZipFile(file) as z:assert len(z.namelist())==2
            if fmt=='docx':
                from docx import Document
                assert 'SECRET' in '\n'.join(p.text for p in Document(file).paragraphs)
                imported,converted=run({'action':'import'},file)
                with fitz.open(converted/'result.pdf') as d:assert 'SECRET' in ''.join(p.get_text() for p in d)
        _,out=run({'action':'redact','pages':'1','rect':[.1,.1,.95,.2]})
        with fitz.open(out/'result.pdf') as d:
            assert 'SECRET' not in d[0].get_text()
            assert 'Keep this content' in d[0].get_text()
            assert d.embfile_count()==0
        _,out=run({'action':'text','pages':'1','rect':[.1,.1,.95,.2],'text':'Replacement 中文','color':'#123456','size':18})
        with fitz.open(out/'result.pdf') as d:
            assert 'SECRET' not in d[0].get_text()
            assert 'Replacement' in d[0].get_text()
            assert '中文' in d[0].get_text()
        for action in ('stamp','watermark','number','shape','annotation','form','crop','flatten','compress','page-stamp'):
            _,out=run({'action':action,'pages':'1','text':'Reviewed {page}/{total}','rect':[.1,.55,.8,.75],'inputs':[str(source)]})
            with fitz.open(out/'result.pdf') as d:
                assert len(d)==2
                if action=='form':assert len(list(d[0].widgets()))==1
                if action=='crop':assert d[0].rect.width<400
                if action=='flatten':assert not list(d[0].annots() or [])
                if action=='number':assert '1/2' in d[0].get_text()
        _,out=run({'action':'organize','mode':'reorder','pages':'2,1,1'})
        with fitz.open(out/'result.pdf') as d:assert len(d)==3 and 'Second' in d[0].get_text()
        for mode in ('delete','rotate','blank','split','merge'):
            result,out=run({'action':'organize','mode':mode,'pages':'1','inputs':[str(source)]})
            if mode=='split':
                with zipfile.ZipFile(out/'result.zip') as z:assert len(z.namelist())==1
            else:
                with fitz.open(out/'result.pdf') as d:
                    assert len(d)=={'delete':1,'rotate':2,'blank':3,'merge':4}[mode]
                    if mode=='rotate':assert d[0].rotation==90
        _,out=run({'action':'outline','toc':[[1,'目录中文',1],[2,'Section',2]]})
        with fitz.open(out/'result.pdf') as d:assert d.get_toc()[0][1]=='目录中文'
        _,encrypted=run({'action':'encrypt','userPassword':'reader-password','ownerPassword':'owner-password','copy':False,'edit':False})
        with fitz.open(encrypted/'result.pdf') as d:
            assert d.needs_pass
            assert not d.authenticate('incorrect')
            assert d.authenticate('reader-password')
            assert not d.permissions&fitz.PDF_PERM_COPY
        try:run({'action':'decrypt','password':'reader-password'},encrypted/'result.pdf');assert False
        except ValueError:pass
        _,out=run({'action':'decrypt','password':'owner-password'},encrypted/'result.pdf')
        with fitz.open(out/'result.pdf') as d:assert not d.needs_pass and 'SECRET' in d[0].get_text()
        _,restricted=run({'action':'encrypt','userPassword':'','ownerPassword':'owner-password','copy':False,'edit':False})
        try:run({'action':'decrypt'},restricted/'result.pdf');assert False
        except ValueError:pass
        try:run({'action':'export','format':'txt'},restricted/'result.pdf');assert False
        except ValueError:pass
        _,out=run({'action':'compare','inputs':[str(source)]})
        assert '逐页文字差异' in (out/'result.html').read_text(encoding='utf-8')
        # OCR a page that contains only an image, not an existing text layer.
        with fitz.open(source) as d:pix=d[0].get_pixmap(dpi=180)
        image=base/'image.png';image.write_bytes(pix.tobytes('png'))
        _,out=run({'action':'import'},image)
        scan=out/'result.pdf'
        with fitz.open(scan) as d:assert not d[0].get_text().strip()
        _,out=run({'action':'ocr','tessdata':str(root/'vendor'/'ocr')},scan)
        with fitz.open(out/'result.pdf') as d:assert 'SECRET' in d[0].get_text().upper()
        # Image stamps keep the picture's aspect ratio by default (no stretching).
        from PIL import Image as PILImage
        picture=base/'stamp-picture.png'
        PILImage.new('RGB',(100,50),(200,30,30)).save(picture)
        def placed(image_request):
            _,out=run({**image_request,'action':'image','asset':str(picture),'rect':[.1,.1,.7,.4]})
            with fitz.open(out/'result.pdf') as d:
                box=fitz.Rect(d[0].get_image_info()[0]['bbox'])
                return box.width/box.height
        kept=placed({'keepRatio':True})
        stretched=placed({'keepRatio':False})
        default=placed({})
        assert abs(kept-2)<.05 and abs(default-2)<.05 and abs(stretched-1.6)<.05,(kept,default,stretched)
        # 0.14 text markup: one annotation per page, one quad per selected line, author and note.
        lines=[[50/400,(80-22)/500,250/400,84/500],[50/400,(180-17)/500,220/400,184/500]]
        for kind,subtype in (('underline','Underline'),('squiggly','Squiggly'),('strikeout','StrikeOut'),('highlight','Highlight')):
            _,out=run({'action':'markup','annotation':kind,'marks':[{'page':1,'rects':lines},{'page':2,'rects':lines[:1]}],'color':'#1d7a57','author':'张三','text':'原文','opacity':.5})
            with fitz.open(out/'result.pdf') as d:
                one,two=d[0],d[1]
                first=[a for a in one.annots() if a.type[1]==subtype]
                assert len(first)==1 and len(first[0].vertices)==8,(kind,len(first))
                assert first[0].info['title']=='张三' and first[0].info['content']=='原文'
                assert 'SECRET' in one.get_textbox(first[0].rect)
                assert len([a for a in two.annots() if a.type[1]==subtype])==1
        # 0.14 redaction by content: presets / terms find exact boxes; only overlapping annotations go.
        sensitive=base/'sensitive.pdf'
        with fitz.open() as doc:
            page=doc.new_page(width=400,height=300)
            page.insert_font(fontname='china',fontbuffer=fitz.Font('cjk').buffer)
            page.insert_text((30,60),'Phone 13812345678 mail a.b@example.com',fontsize=12)
            page.insert_text((30,120),'客户：王小明 身份证 11010519491231002X',fontsize=12,fontname='china')
            page.insert_text((30,200),'Public line stays',fontsize=12)
            page.add_highlight_annot(fitz.Rect(70,48,150,64))
            page.add_text_annot((300,200),'unrelated note')
            doc.save(sensitive)
        found,_=run({'action':'redact-find','presets':['mobile','email','idcard'],'terms':['王小明']},sensitive)
        hits=found['data']['hits']
        assert sorted(h['text'] for h in hits)==sorted(['13812345678','a.b@example.com','11010519491231002X','王小明']),hits
        assert all(len(h['rects'])==1 and 0<=h['rects'][0][0]<h['rects'][0][2]<=1 for h in hits)
        try:run({'action':'redact-find'},sensitive);assert False
        except ValueError:pass
        try:run({'action':'redact-find','patterns':['(']},sensitive);assert False
        except ValueError:pass
        _,out=run({'action':'redact','marks':[{'page':h['page'],'rects':h['rects']} for h in hits]},sensitive)
        with fitz.open(out/'result.pdf') as d:
            page=d[0];text=page.get_text()
            for secret in ('13812345678','example.com','1949','王小明'):assert secret not in text,(secret,text)
            assert 'Phone' in text and 'Public line stays' in text and '客户' in text
            kinds=[a.type[1] for a in page.annots()]
            assert 'Highlight' not in kinds and 'Text' in kinds,kinds
        # The framed-region path no longer flattens annotations elsewhere on the page.
        _,out=run({'action':'redact','pages':'1','rect':[0,.6,.7,.7]},sensitive)
        with fitz.open(out/'result.pdf') as d:
            page=d[0]
            assert 'Public line stays' not in page.get_text()
            assert sorted(a.type[1] for a in page.annots())==['Highlight','Text']
        # 0.14 stamp templates, Bates numbers and odd / even pages.
        four=base/'four.pdf'
        with fitz.open() as doc:
            for n in range(4):doc.new_page(width=300,height=300)
            doc.set_metadata({'title':'合同'});doc.save(four)
        _,out=run({'action':'number','text':'{bates} · {page}/{total} · {title} · {file}','subset':'even','batesPrefix':'QY-','batesStart':7,'batesDigits':4,'fileName':'a.pdf','rect':[.05,.85,.95,.97],'size':10},four)
        with fitz.open(out/'result.pdf') as d:
            texts=[d[i].get_text().strip() for i in range(4)]
        assert texts[0]=='' and texts[2]=='',texts
        assert 'QY-0007 · 2/4 · 合同 · a.pdf' in texts[1] and 'QY-0008 · 4/4' in texts[3],texts
        try:run({'action':'stamp','text':'x','subset':'even','pages':'1'},four);assert False
        except ValueError:pass
        # 0.14 compression profiles: a 600 DPI photo is downsampled; smaller profile → smaller image.
        photo=base/'photo.pdf';noise=base/'noise.png'
        PILImage.effect_noise((1200,1200),60).convert('RGB').save(noise)
        with fitz.open() as doc:
            doc.new_page(width=144,height=144).insert_image(fitz.Rect(0,0,144,144),filename=str(noise));doc.save(photo)
        widths={}
        for profile in ('quality','balanced','small'):
            _,out=run({'action':'compress','profile':profile},photo)
            with fitz.open(out/'result.pdf') as d:widths[profile]=d[0].get_images()[0][2]
        assert 1200>widths['quality']>=widths['balanced']>widths['small'],widths  # MuPDF subsamples by integer factors
        # 0.14 drawn shapes and deleting annotations by xref.
        made={}
        for shape,request in (('rect',{'rect':[.1,.1,.4,.3],'fill':'#ffee00'}),('ellipse',{'rect':[.5,.1,.9,.3]}),('cloud',{'rect':[.1,.5,.4,.7]}),('line',{'points':[[.1,.8],[.9,.8]]}),('arrow',{'points':[[.1,.9],[.9,.85]]})):
            _,out=run({'action':'draw','shape':shape,'page':1,'color':'#1d7a57','width':3,'author':'李四',**request})
            with fitz.open(out/'result.pdf') as d:
                one=d[0];annot=[a for a in one.annots() if a.info['title']=='李四'][0]
                made[shape]=(annot.type[1],annot.border.get('width'),annot.line_ends if shape in ('line','arrow') else None,annot.colors.get('fill'))
        assert made['rect'][0]=='Square' and made['rect'][3] and made['ellipse'][0]=='Circle' and made['cloud'][0]=='Square',made
        assert made['line'][0]=='Line' and made['arrow'][2][1]==fitz.PDF_ANNOT_LE_OPEN_ARROW and made['rect'][1]==3,made
        try:run({'action':'draw','shape':'line','page':1,'points':[[.1,.1],[.1,.1]]});assert False
        except ValueError:pass
        _,out=run({'action':'draw','shape':'rect','page':1,'rect':[.1,.1,.4,.3]},source)
        with fitz.open(out/'result.pdf') as d:
            xrefs=[(a.xref,a.type[1]) for a in d[0].annots()]
        target=[x for x,t in xrefs if t=='Square'][0]
        _,cleaned=run({'action':'annot-delete','items':[{'page':1,'xref':target}]},out/'result.pdf')
        with fitz.open(cleaned/'result.pdf') as d:
            assert [a.type[1] for a in d[0].annots()]==[t for x,t in xrefs if x!=target]
        gone,_=run({'action':'annot-delete','items':[{'page':1,'xref':999999}]},source);assert gone.get('unchanged')
        for bad in ({'annotation':'circle','marks':[{'page':1,'rects':lines}]},{'annotation':'underline','marks':[]},{'annotation':'underline','marks':[{'page':1,'rects':[]}]}):
            try:run({'action':'markup',**bad});assert False,bad
            except ValueError:pass
        print(f'BACKEND PASS: {counter} jobs, exports, edits, encryption, redaction, OCR, outline, pages and comparison')


if __name__=='__main__':check()
